#!/usr/bin/env python3
"""
extract_translations.py — Target 08, מדידת נפח מוצק: translation export (he → ar).

Reads every learner-facing Hebrew string in the build and writes, into translation/:
  methodica-science-volume-solid-01-he-ar.xlf   XLIFF 1.2, one <trans-unit> per string, empty <target>
  methodica-science-volume-solid-01-he-ar.xlsx  the same units as a spreadsheet for the translator
  translation-screen-manifest.csv               one row per screen

Read-only against the build. Run from the repo root:
    python3 translation/tools/extract_translations.py

What counts as a unit
  HTML  The highest element whose content is Hebrew text plus inline markup only (span, b, strong,
        br, …) — its inner HTML is the source, so a sentence split across styling runs stays one
        unit. Hebrew text that sits directly beside block children becomes its own unit. Every
        attribute holding Hebrew (alt, aria-label, title, placeholder, data-*) is a unit too.
        Comments, <script> and <style> are skipped.
  JS    Every string literal (', ", `) holding Hebrew, comments skipped. Concatenated pieces stay
        separate units; the note says so.

IDs are 08-PP-SS-NNN: part, screen (data-screen; 99 = part-wide chrome / JS not tied to a screen;
part 00 = the module root and the shared unit-js layer), running number.
"""
import csv
import html
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from bs4 import BeautifulSoup, Comment, NavigableString, Tag

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "translation"
SLUG = "methodica-science-volume-solid-01"
TARGET = "08"
PARTS = ["01", "02", "03", "04", "05", "06", "07"]

HEB = re.compile(r"[֐-׿]")
WS = re.compile(r"\s+")
INLINE = {"span", "b", "strong", "em", "i", "u", "sup", "sub", "br", "wbr", "small", "a",
          "bdi", "bdo", "mark", "s", "q", "abbr", "code", "kbd"}
SKIP_TAGS = {"script", "style", "noscript"}
SKIP_ATTRS = {"class", "id", "style", "src", "href"}

ROLE_BY_TAG = {
    "h1": "heading", "h2": "heading", "h3": "subheading", "h4": "subheading", "p": "paragraph",
    "li": "list-item", "button": "button-label", "label": "label", "option": "option",
    "title": "page-title", "th": "table-header", "td": "table-cell", "figcaption": "caption",
    "legend": "label", "text": "svg-text", "tspan": "svg-text",
}
ROLE_BY_ATTR = {"alt": "image-alt", "aria-label": "aria-label", "title": "tooltip",
                "placeholder": "placeholder", "aria-roledescription": "aria-roledescription"}

TYPE_HINTS = [  # (substring in the section's HTML, screen type) — first hit wins
    ("comic-", "Comic panel screen"), ("peak-opt", "Peak question (SingleChoice)"),
    ("ddq", "DragAndDropQuestion"), ("dq-sentence", "DropdownQuestion"),
    ("mcq", "MultipleChoiceQuestion"), ("meas-", "Measurement applet"),
    ("disp-", "Displacement applet"), ("flood-", "Overflow applet"),
    ("scq-opt", "SingleChoiceQuestion"), ("guess-options", "Guess question"),
    ("transition-", "TransitionScreen"), ("path-choice", "TwoOptionSelection"),
]


def norm(s):
    return WS.sub(" ", s).strip()


def has_heb(s):
    return bool(s) and bool(HEB.search(s))


def inline_only(el):
    return all(d.name in INLINE for d in el.descendants if isinstance(d, Tag))


def selector(el):
    s = el.name
    if el.get("id"):
        s += "#" + el["id"]
    cls = el.get("class") or []
    if cls:
        s += "." + ".".join(cls[:2])
    return s


class Unit:
    def __init__(self, part, screen, source, role, loc, element="", context=""):
        self.part, self.screen, self.source = part, screen, source
        self.role, self.loc, self.element, self.context = role, loc, element, context
        self.id = ""


# ── HTML ──────────────────────────────────────────────────────────────────────────────────────
def html_units(path, rel, part, screens_out):
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    for c in soup.find_all(string=lambda s: isinstance(s, Comment)):
        c.extract()
    for t in soup.find_all(SKIP_TAGS):
        t.decompose()

    dom_to_screen = {}
    sections = soup.select("section.screen[data-screen]")
    for sec in sections:
        dom_to_screen[sec.get("id", "")] = int(sec["data-screen"])

    def screen_of(el):
        for a in [el] + list(el.parents):
            if not isinstance(a, Tag):
                continue
            if a.name == "section" and a.get("data-screen") is not None and "screen" in (a.get("class") or []):
                return int(a["data-screen"])
            m = re.match(r"s(\d+)-", a.get("id", "") or "")
            if m and ("s" + m.group(1)) in dom_to_screen:
                return dom_to_screen["s" + m.group(1)]
        return 99

    units = []

    def emit_attrs(el):
        for k, v in el.attrs.items():
            if k in SKIP_ATTRS or k.startswith("on"):
                continue
            v = " ".join(v) if isinstance(v, list) else v
            if has_heb(v):
                role = ROLE_BY_ATTR.get(k, "attribute:" + k)
                units.append(Unit(part, screen_of(el), norm(v), role, f"{rel}:{el.sourceline}",
                                  selector(el) + f"[{k}]"))

    def visit(el):
        emit_attrs(el)
        if not has_heb(el.get_text()):
            for d in el.find_all(True):
                emit_attrs(d)
            return
        if el.name not in ("html", "body", "head") and inline_only(el):
            role = ROLE_BY_TAG.get(el.name, "text")
            units.append(Unit(part, screen_of(el), norm(el.decode_contents()), role,
                              f"{rel}:{el.sourceline}", selector(el)))
            for d in el.find_all(True):
                emit_attrs(d)
            return
        for ch in el.children:
            if isinstance(ch, Tag):
                visit(ch)
            elif isinstance(ch, NavigableString) and has_heb(str(ch)):
                units.append(Unit(part, screen_of(el), norm(str(ch)), "text",
                                  f"{rel}:{el.sourceline}", selector(el) + " (loose text)"))

    visit(soup)

    for sec in sections:
        n = int(sec["data-screen"])
        head = sec.select_one("h1, h2, .exp-title, .transition-title, .scq-question-text")
        title = norm(head.get_text()) if head else ""
        raw = str(sec)
        stype = next((t for k, t in TYPE_HINTS if k in raw), "other")
        screens_out.append((part, n, sec.get("id", ""), title, stype, rel))
    return units


# ── JS ────────────────────────────────────────────────────────────────────────────────────────
REGEX_PRE = set("(,=:[!&|?{};+-*%<>~^")
KEY_SCREEN = re.compile(r"^\s{2}'?s(\d{1,2})'?\s*:")
KEY_NUM = re.compile(r"^\s{2}(\d{1,2})\s*:\s*[\[{]")
TOPLEVEL = re.compile(r"^\S")
KEY_BEFORE = re.compile(r"([A-Za-z_$][\w$]*)\s*:\s*$")


def js_literals(src):
    """Yield (start_index, text) for every string literal, skipping comments and regex literals."""
    i, n, prev = 0, len(src), ""
    while i < n:
        c = src[i]
        if c == "/" and i + 1 < n and src[i + 1] == "/":
            j = src.find("\n", i)
            i = n if j < 0 else j
            continue
        if c == "/" and i + 1 < n and src[i + 1] == "*":
            j = src.find("*/", i + 2)
            i = n if j < 0 else j + 2
            continue
        if c == "/" and (prev in REGEX_PRE or prev == ""):
            j = i + 1
            in_cls = False
            while j < n and src[j] != "\n":
                if src[j] == "\\":
                    j += 2
                    continue
                if src[j] == "[":
                    in_cls = True
                elif src[j] == "]":
                    in_cls = False
                elif src[j] == "/" and not in_cls:
                    break
                j += 1
            i, prev = j + 1, "x"
            continue
        if c in "'\"`":
            j, buf = i + 1, []
            while j < n and src[j] != c:
                if src[j] == "\\" and j + 1 < n:
                    buf.append(src[j:j + 2])
                    j += 2
                    continue
                if c == "`" and src[j] == "$" and j + 1 < n and src[j + 1] == "{":
                    depth, k = 1, j + 2
                    while k < n and depth:
                        depth += {"{": 1, "}": -1}.get(src[k], 0)
                        k += 1
                    buf.append(src[j:k])
                    j = k
                    continue
                buf.append(src[j])
                j += 1
            yield i, "".join(buf)
            i, prev = j + 1, "x"
            continue
        if not c.isspace():
            prev = c
        i += 1


def unescape_js(s):
    s = s.replace("\\n", "\n").replace("\\'", "'").replace('\\"', '"').replace("\\`", "`")
    s = re.sub(r"\\u([0-9A-Fa-f]{4})", lambda m: chr(int(m.group(1), 16)), s)
    return s.replace("\\\\", "\\")


def js_units(path, rel, part, peak=False):
    src = path.read_text(encoding="utf-8")
    lines = src.split("\n")
    line_start = [0]
    for ln in lines:
        line_start.append(line_start[-1] + len(ln) + 1)

    # screen per line: inside a `  sNN: {` (or, in Peak parts, `  N: {`) block of a top-level object
    screen_at, cur = [], 99
    for ln in lines:
        if TOPLEVEL.match(ln):
            cur = 99
        m = KEY_SCREEN.match(ln) or (KEY_NUM.match(ln) if peak else None)
        if m:
            cur = int(m.group(1))
        screen_at.append(cur)

    import bisect
    units, last_line = [], -1
    for pos, text in js_literals(src):
        if not has_heb(text):
            continue
        lno = bisect.bisect_right(line_start, pos) - 1
        before = src[line_start[lno]:pos]
        km = KEY_BEFORE.search(before)
        key = km.group(1) if km else ""
        role = {"alt": "image-alt", "title": "title", "text": "text", "body": "feedback-body",
                "placeholder": "placeholder"}.get(key, "js-string" + (":" + key if key else ""))
        tm = re.search(r"type:\s*'(\w+)'", lines[lno])
        if tm and key == "text":
            role = "comic-" + tm.group(1)
        concat = before.rstrip().endswith("+") or src[pos:pos + len(text) + 40].split("\n")[0].count("+") > 0
        ctx = norm(lines[lno])[:160]
        u = Unit(part, screen_at[lno], norm(unescape_js(text)) if "\n" not in unescape_js(text)
                 else unescape_js(text).strip(), role, f"{rel}:{lno + 1}", "", ctx)
        if concat:
            u.context = "(part of a + concatenation) " + ctx
        units.append(u)
    return units


# ── main ──────────────────────────────────────────────────────────────────────────────────────
def main():
    units, screens = [], []
    # Part 00: module root + shared unit-js layer
    units += html_units(ROOT / "index.html", "index.html", "00", [])
    for f in sorted((ROOT / "unit-js").glob("*.js")):
        units += js_units(f, f"unit-js/{f.name}", "00")
    for p in PARTS:
        d = f"{SLUG}-{p}"
        units += html_units(ROOT / d / "index.html", f"{d}/index.html", p, screens)
        for f in sorted((ROOT / d / "js").glob("*.js")):
            units += js_units(f, f"{d}/js/{f.name}", p, peak=p in ("06", "07"))
    for u in units:
        if u.part == "00":
            u.screen = 0 if u.loc.startswith("index.html") else 99

    # stable order: part, screen, file order; number per screen
    order = {k: i for i, k in enumerate(dict.fromkeys(u.loc.split(":")[0] for u in units))}
    units.sort(key=lambda u: (u.part, u.screen, order[u.loc.split(":")[0]], int(u.loc.split(":")[1] or 0)))
    seq = {}
    for u in units:
        k = (u.part, u.screen)
        seq[k] = seq.get(k, 0) + 1
        u.id = f"{TARGET}-{u.part}-{u.screen:02d}-{seq[k]:03d}"

    sinfo = {(p, n): (sid, t, ty) for p, n, sid, t, ty, _ in screens}

    def screen_title(u):
        if u.screen == 99:
            return "Part-wide chrome / runtime strings (not tied to one screen)"
        if u.part == "00":
            return "Module root redirect page"
        return sinfo.get((u.part, u.screen), ("", "", ""))[1]

    def screen_type(u):
        if u.screen == 99 or u.part == "00":
            return "chrome"
        return sinfo.get((u.part, u.screen), ("", "", ""))[2]

    OUT.mkdir(exist_ok=True)
    # XLIFF
    xl = ['<?xml version="1.0" encoding="UTF-8"?>',
          '<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">',
          f'  <file original="{SLUG}" source-language="he" target-language="ar" datatype="html" '
          'tool-id="claude-code-translation-export">',
          "    <header>",
          "      <note>Target/learning-objective 08 — מדידת נפח מוצק (Measuring the volume of solids), "
          "720-series Hebrew e-learning module. Parts 01-07 plus pseudo-part 00 (module root and the shared "
          "unit-js layer). Sources holding HTML tags keep them: translate the text, keep the tags. "
          "No production files were modified to produce this export.</note>",
          "    </header>", "    <body>"]
    for u in units:
        xl.append(f'      <trans-unit id="{u.id}">')
        xl.append(f"        <source>{escape(u.source)}</source>")
        xl.append("        <target></target>")
        xl.append(f'        <note from="screen-id">{TARGET}-{u.part}-{u.screen:02d}</note>')
        t = screen_title(u)
        if t:
            xl.append(f'        <note from="screen-title">{escape(t)}</note>')
        xl.append(f'        <note from="screen-type">{escape(screen_type(u))}</note>')
        xl.append(f'        <note from="element-role">{escape(u.role)}</note>')
        if u.element:
            xl.append(f'        <note from="element">{escape(u.element)}</note>')
        if u.context:
            xl.append(f'        <note from="js-context">{escape(u.context)}</note>')
        xl.append(f'        <note from="source-location">{escape(u.loc)}</note>')
        xl.append("      </trans-unit>")
    xl += ["    </body>", "  </file>", "</xliff>", ""]
    (OUT / f"{SLUG}-he-ar.xlf").write_text("\n".join(xl), encoding="utf-8")

    # manifest
    with open(OUT / "translation-screen-manifest.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["screen_id", "dom_id", "screen_title", "screen_type", "source_file", "units"])
        cnt = {}
        for u in units:
            cnt[(u.part, u.screen)] = cnt.get((u.part, u.screen), 0) + 1
        for p, n, sid, t, ty, rel in sorted(screens):
            w.writerow([f"{TARGET}-{p}-{n:02d}", sid, t, ty, rel, cnt.get((p, n), 0)])

    # xlsx
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "תרגום"
    ws.sheet_view.rightToLeft = True
    hdr = ["מזהה", "חלק", "מסך", "כותרת המסך", "תפקיד הטקסט", "טקסט מקור (עברית)", "תרגום (ערבית)",
           "הערות", "מיקום בקוד"]
    ws.append(hdr)
    for u in units:
        notes = []
        if "<" in u.source:
            notes.append("יש תגיות HTML — לתרגם את הטקסט ולהשאיר את התגיות")
        if u.context.startswith("(part of"):
            notes.append("חלק ממשפט שמחובר בקוד — לתרגם יחד עם השורות הסמוכות")
        ws.append([u.id, u.part, f"{u.screen:02d}", screen_title(u), u.role, u.source, "",
                   "; ".join(notes), u.loc])
    widths = [15, 6, 6, 34, 16, 60, 60, 30, 40]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6FA3")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[6].fill = PatternFill("solid", fgColor="FFF7D6")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(units) + 1}"
    ms = wb.create_sheet("מסכים")
    ms.sheet_view.rightToLeft = True
    ms.append(["מזהה מסך", "DOM", "כותרת", "סוג", "קובץ", "יחידות"])
    with open(OUT / "translation-screen-manifest.csv", encoding="utf-8") as fh:
        for r in list(csv.reader(fh))[1:]:
            ms.append(r)
    for i, wdt in enumerate([12, 8, 50, 28, 45, 8]):
        ms.column_dimensions[chr(65 + i)].width = wdt
    wb.save(OUT / f"{SLUG}-he-ar.xlsx")

    by_part = {}
    for u in units:
        by_part[u.part] = by_part.get(u.part, 0) + 1
    ids = [u.id for u in units]
    print(f"units={len(units)} unique_ids={len(set(ids))} by_part={dict(sorted(by_part.items()))}")
    print(f"screens={len(screens)} screen99={sum(1 for u in units if u.screen == 99)}")


if __name__ == "__main__":
    sys.exit(main())
